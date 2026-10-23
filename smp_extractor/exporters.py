from __future__ import annotations

import csv
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import ExtractedTable, SchoolResult

TOOL_VERSION = "0.1.0"


def _slug(value: str, fallback: str = "table") -> str:
    value = re.sub(r"[^A-Za-z0-9_-]+", "_", value.strip())
    value = re.sub(r"_+", "_", value).strip("_-")
    return value[:48] or fallback


def _unique_sheet_name(name: str, used: set[str]) -> str:
    base = re.sub(r"[\\/*?:\[\]]", "_", name).strip("' ") or "Table"
    base = base[:31]
    candidate = base
    counter = 2
    while candidate.casefold() in used:
        suffix = f"_{counter}"
        candidate = f"{base[:31 - len(suffix)]}{suffix}"
        counter += 1
    used.add(candidate.casefold())
    return candidate


def _write_excel_text(cell: Any, value: Any) -> None:
    """Store portal content as text, never as an executable Excel formula."""
    cell.value = "" if value is None else str(value)
    cell.data_type = "s"
    cell.number_format = "@"


def _safe_source_url(value: str) -> str:
    """Mask credential-like query parameters before writing provenance metadata."""
    parsed = urlsplit(value)
    safe_query = []
    for key, query_value in parse_qsl(parsed.query, keep_blank_values=True):
        normalized_key = re.sub(r"[^a-z0-9]", "", key.casefold())
        is_sensitive = (
            any(token in normalized_key for token in ("token", "auth", "session", "password", "secret", "csrf", "xsrf", "apikey"))
            or normalized_key in {"key", "sid"}
        )
        safe_query.append((key, "[redacted]" if is_sensitive else query_value))
    return urlunsplit(
        (parsed.scheme, parsed.netloc, parsed.path, urlencode(safe_query, doseq=True), "")
    )


def _safe_csv_value(value: Any) -> str:
    """Neutralize formula-like text when the CSV is opened in a spreadsheet app.

    XLSX contains the original string as an explicit text cell. CSV has no cell
    type, so an apostrophe is prefixed to values Excel could interpret as a
    formula. Values in ordinary text editors remain easy to identify.
    """
    text = "" if value is None else str(value)
    probe = text.lstrip(" \t\r\n")
    if probe.startswith(("=", "+", "-", "@")) or text.startswith(("\t", "\r", "\n")):
        return "'" + text
    return text


def _write_csv(table: ExtractedTable, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle, quoting=csv.QUOTE_MINIMAL)
        writer.writerow([_safe_csv_value(header) for header in table.headers])
        for row in table.rows:
            normalized = list(row[: len(table.headers)])
            if len(normalized) < len(table.headers):
                normalized.extend([""] * (len(table.headers) - len(normalized)))
            writer.writerow([_safe_csv_value(value) for value in normalized])


def _style_header(worksheet: Any, row_number: int = 1) -> None:
    fill = PatternFill(fill_type="solid", fgColor="17365D")
    font = Font(color="FFFFFF", bold=True)
    for cell in worksheet[row_number]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(vertical="top", wrap_text=True)


def _fit_columns(worksheet: Any, max_rows: int = 500) -> None:
    if worksheet.max_column == 0:
        return
    for column_index in range(1, worksheet.max_column + 1):
        longest = 0
        for row_index in range(1, min(worksheet.max_row, max_rows) + 1):
            value = worksheet.cell(row=row_index, column=column_index).value
            longest = max(longest, len(str(value or "")))
        worksheet.column_dimensions[get_column_letter(column_index)].width = min(
            max(longest + 2, 12), 55
        )


def _safe_sheet(worksheet: Any, headers: list[str], rows: list[list[str]]) -> None:
    effective_headers = headers or ["No columns detected"]
    for column, header in enumerate(effective_headers, start=1):
        _write_excel_text(worksheet.cell(row=1, column=column), header)
    _style_header(worksheet)
    worksheet.freeze_panes = "A2"

    for row_index, row in enumerate(rows, start=2):
        for column_index, value in enumerate(row[: len(effective_headers)], start=1):
            _write_excel_text(worksheet.cell(row=row_index, column=column_index), value)
        for column_index in range(len(row) + 1, len(effective_headers) + 1):
            _write_excel_text(worksheet.cell(row=row_index, column=column_index), "")

    if worksheet.max_row >= 1 and worksheet.max_column >= 1:
        worksheet.auto_filter.ref = worksheet.dimensions
    _fit_columns(worksheet)


def _output_directory(output_root: Path) -> tuple[Path, str]:
    output_root.mkdir(parents=True, exist_ok=True)
    now = datetime.now().astimezone()
    run_id = now.strftime("%Y%m%d_%H%M%S")
    output_dir = output_root / run_id
    suffix = 2
    while output_dir.exists():
        output_dir = output_root / f"{run_id}_{suffix:02d}"
        suffix += 1
    output_dir.mkdir(parents=True)
    return output_dir, datetime.now(timezone.utc).isoformat(timespec="seconds")


def write_export_bundle(
    results: list[SchoolResult],
    base_url: str,
    output_root: str | Path = "exports",
) -> Path:
    """Write one workbook, a CSV per table, and a metadata-only manifest."""
    output_dir, generated_at = _output_directory(Path(output_root))
    csv_dir = output_dir / "csv"
    csv_dir.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary_headers = [
        "EMIS code",
        "School page title",
        "Status",
        "Table #",
        "Table title",
        "Rows captured",
        "Columns",
        "Workbook sheet",
        "CSV file",
        "Source URLs",
        "Capture contexts",
        "Warnings / error",
    ]
    for column, header in enumerate(summary_headers, start=1):
        _write_excel_text(summary.cell(row=1, column=column), header)
    _style_header(summary)
    summary.freeze_panes = "A2"

    used_sheet_names = {"summary"}
    manifest: dict[str, Any] = {
        "tool": "smp-table-extractor",
        "version": TOOL_VERSION,
        "generated_at_utc": generated_at,
        "base_url": base_url,
        "requested_codes": [result.emis_code for result in results],
        "results": [],
    }
    sequence = 0

    for result in results:
        result_manifest: dict[str, Any] = {
            "emis_code": result.emis_code,
            "url": result.url,
            "page_title": result.page_title,
            "status": result.status,
            "table_count": result.table_count,
            "row_count": result.row_count,
            "warnings": list(result.warnings),
            "error": result.error,
            "tables": [],
        }

        if not result.tables:
            summary_row = [
                result.emis_code,
                result.page_title,
                result.status,
                "",
                "",
                0,
                0,
                "",
                "",
                result.url,
                "",
                "; ".join(result.warnings + ([result.error] if result.error else [])),
            ]
            _append_summary_row(summary, summary_row)

        for table in result.tables:
            sequence += 1
            sheet_base = f"{result.emis_code}_{sequence:02d}_{_slug(table.title)}"
            sheet_name = _unique_sheet_name(sheet_base, used_sheet_names)
            worksheet = workbook.create_sheet(sheet_name)
            _safe_sheet(worksheet, table.headers, table.rows)

            csv_name = f"{result.emis_code}_{sequence:02d}_{_slug(table.title)}.csv"
            csv_path = csv_dir / csv_name
            _write_csv(table, csv_path)
            relative_csv = str(csv_path.relative_to(output_dir)).replace("\\", "/")
            safe_source_urls = [_safe_source_url(url) for url in table.source_urls]

            summary_row = [
                result.emis_code,
                result.page_title,
                result.status,
                table.table_index + 1,
                table.title,
                table.row_count,
                table.column_count,
                sheet_name,
                relative_csv,
                " | ".join(safe_source_urls),
                " | ".join(table.capture_contexts),
                "; ".join(result.warnings + ([result.error] if result.error else [])),
            ]
            _append_summary_row(summary, summary_row)

            result_manifest["tables"].append(
                {
                    "table_number": table.table_index + 1,
                    "title": table.title,
                    "headers": list(table.headers),
                    "row_count": table.row_count,
                    "column_count": table.column_count,
                    "workbook_sheet": sheet_name,
                    "csv_file": relative_csv,
                    "source_urls": safe_source_urls,
                    "capture_contexts": list(table.capture_contexts),
                }
            )
        manifest["results"].append(result_manifest)

    _fit_columns(summary)
    workbook.properties.creator = "SMP table extractor"
    workbook.properties.title = "SMP school table export"
    workbook.save(output_dir / "smp_tables.xlsx")

    manifest["output_files"] = [
        "smp_tables.xlsx",
        "run_manifest.json",
        *[
            str(path.relative_to(output_dir)).replace("\\", "/")
            for path in sorted(csv_dir.glob("*.csv"))
        ],
    ]
    with (output_dir / "run_manifest.json").open("w", encoding="utf-8") as handle:
        json.dump(manifest, handle, ensure_ascii=False, indent=2)
        handle.write("\n")

    return output_dir


def _append_summary_row(worksheet: Any, values: list[Any]) -> None:
    row_index = worksheet.max_row + 1
    for column_index, value in enumerate(values, start=1):
        cell = worksheet.cell(row=row_index, column=column_index)
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            cell.value = value
        else:
            _write_excel_text(cell, value)
