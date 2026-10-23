import csv
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from smp_extractor.exporters import _safe_source_url, write_export_bundle
from smp_extractor.models import ExtractedTable, SchoolResult


class ExporterTests(unittest.TestCase):
    def test_writes_workbook_csv_and_manifest_safely(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            table = ExtractedTable(
                emis_code="32230155",
                table_index=0,
                table_ref="school-info",
                title="School information",
                headers=["Field", "Value"],
                rows=[
                    ["School name", "Example School"],
                    ["Formula-like text", "=HYPERLINK(\"https://bad.invalid\",\"click\")"],
                ],
                source_urls=["https://smp26.pesrp.edu.pk/school-detail/32230155"],
                capture_contexts=["initial page state"],
            )
            result = SchoolResult(
                emis_code="32230155",
                url="https://smp26.pesrp.edu.pk/school-detail/32230155",
                status="ok",
                page_title="Example School",
                tables=[table],
            )

            output_dir = write_export_bundle(
                [result], "https://smp26.pesrp.edu.pk", Path(temp_dir) / "exports"
            )

            workbook_path = output_dir / "smp_tables.xlsx"
            self.assertTrue(workbook_path.exists())
            workbook = load_workbook(workbook_path, data_only=False)
            self.assertEqual(workbook.sheetnames[0], "Summary")
            data_sheet = workbook[workbook.sheetnames[1]]
            self.assertEqual(data_sheet.cell(1, 1).value, "Field")
            formula_cell = data_sheet.cell(3, 2)
            self.assertEqual(
                formula_cell.value,
                '=HYPERLINK("https://bad.invalid","click")',
            )
            self.assertEqual(formula_cell.data_type, "s")
            self.assertEqual(workbook["Summary"].cell(2, 1).value, "32230155")

            csv_path = output_dir / "csv" / "32230155_01_School_information.csv"
            self.assertTrue(csv_path.exists())
            with csv_path.open("r", encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.reader(handle))
            self.assertEqual(rows[0], ["Field", "Value"])
            self.assertEqual(rows[1], ["School name", "Example School"])
            self.assertTrue(rows[2][1].startswith("'=HYPERLINK"))

            manifest_path = output_dir / "run_manifest.json"
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            self.assertEqual(manifest["requested_codes"], ["32230155"])
            self.assertEqual(manifest["results"][0]["tables"][0]["row_count"], 2)
            self.assertNotIn("password", manifest)

    def test_source_url_masks_secret_query_parameters(self):
        safe = _safe_source_url(
            "https://example.test/school-detail/1?page=2&session_token=secret&district=North"
        )
        self.assertIn("page=2", safe)
        self.assertIn("district=North", safe)
        self.assertNotIn("secret", safe)
        self.assertIn("redacted", safe)

    def test_empty_school_is_still_reported_in_workbook_and_manifest(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            result = SchoolResult(
                emis_code="32230183",
                url="https://smp26.pesrp.edu.pk/school-detail/32230183",
                status="no_tables",
                warnings=["No HTML tables were found."],
            )
            output_dir = write_export_bundle(
                [result], "https://smp26.pesrp.edu.pk", Path(temp_dir) / "exports"
            )
            workbook = load_workbook(output_dir / "smp_tables.xlsx", data_only=True)
            summary = workbook["Summary"]
            self.assertEqual(summary.cell(2, 1).value, "32230183")
            self.assertEqual(summary.cell(2, 3).value, "no_tables")
            manifest = json.loads((output_dir / "run_manifest.json").read_text())
            self.assertEqual(manifest["results"][0]["table_count"], 0)


if __name__ == "__main__":
    unittest.main()
